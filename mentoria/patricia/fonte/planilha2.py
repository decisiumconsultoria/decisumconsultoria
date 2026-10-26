# Planilha simplificada: dias e conteúdo já preenchidos; a aluna marca ✓ e anota questões.
import sys, datetime as dt, plano
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference

OUT = sys.argv[1]; F = "Arial"
TEAL = "1D6B62"
fill = lambda c: PatternFill("solid", fgColor=c)
INP, WEEKA, TODAYF, GREEN, YEL, RED, GREY = fill("FFF4C2"), fill("F4F7F6"), fill("D6E9FF"), fill("CDEBD9"), fill("FFF0C2"), fill("F6D3CF"), fill("ECEFEE")
thin = Side(style="thin", color="D3DBD8"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True); small = Font(name=F, size=8.5, color="66736F")
H = Font(name=F, bold=True, color="FFFFFF", size=10); HF = fill(TEAL)
title = Font(name=F, size=18, bold=True, color=TEAL); sub = Font(name=F, size=10, color="5A6863")
C = Alignment(horizontal="center", vertical="center", wrap_text=True); W = Alignment(vertical="center", wrap_text=True)
WEEKDAYS = {0:"Segunda",2:"Quarta",3:"Quinta",4:"Sexta",6:"Domingo"}
DISC = ["Medicina Legal","Português","Direito"]

wb = Workbook()
def sheet(name, first=False, tab=None):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name; ws.sheet_view.showGridLines = False
    if tab: ws.sheet_properties.tabColor = tab
    return ws
def head(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=i, value=c); x.font = H; x.fill = HF; x.alignment = C; x.border = box
        ws.column_dimensions[x.column_letter].width = w
    ws.row_dimensions[row].height = 30
def put(ws, r, c, v=None, font=base, f=None, al=C, fmt=None):
    x = ws.cell(row=r, column=c, value=v); x.font = font; x.alignment = al; x.border = box
    if f: x.fill = f
    if fmt: x.number_format = fmt
    return x
def dv(ws, rng, formula=None, kind="list", **kw):
    d = DataValidation(type=kind, formula1=formula, allow_blank=True, **kw); ws.add_data_validation(d)
    for part in rng.split(): d.add(part)
    return d
def top(ws, t, s):
    ws["A1"] = t; ws["A1"].font = title; ws["A2"] = s; ws["A2"].font = sub; ws.row_dimensions[1].height = 28

# ---------------- plano diário ----------------
def blocks(n, wd):
    ml, pt, di = plano.week_short(n)
    if n in plano.RECOVERY: return ("Revisão","Revisão de 30 dias e atrasos"), ("Livre","Livre (opcional)")
    if n == 26: return ("Revisão","Revisões leves"), ("Livre","Livre (descanso)")
    sim_sun = n in (21,23,25)
    corr_mon = n in (22,24)
    if n <= 20:
        ptT = (f"Teoria: {pt}" if n <= 18 else pt); ptQ = (f"Questões: {pt}" if n <= 18 else pt)
        diT = (f"Teoria e lei seca: {di}" if n <= 19 else di); diQ = (f"Questões: {di}" if n <= 19 else di)
        table = {0: (("Medicina Legal", f"Teoria: {ml}"), ("Medicina Legal", f"Questões: {ml}")),
                 2: (("Português", ptT), ("Direito", diT)),
                 3: (("Medicina Legal", f"Teoria (continuação): {ml}"), ("Português", ptQ)),
                 4: (("Medicina Legal", f"Questões e flashcards: {ml}"), ("Direito", diQ)),
                 6: (("Português", "Interpretação de textos FGV"), ("Revisão", "Revisão da semana + painel"))}
        if n == 13 and wd == 6: return ("Simulado","Mini-simulado diagnóstico (40 questões)"), ("Simulado","Mini-simulado (continuação)")
    else:
        table = {0: (("Revisão","Correção do simulado") if corr_mon else ("Medicina Legal", ml), ("Medicina Legal", ml)),
                 2: (("Português", pt), ("Direito", di)),
                 3: (("Medicina Legal", "Teoria dos tópicos com mais erros"), ("Português", pt)),
                 4: (("Medicina Legal", ml), ("Direito", di)),
                 6: (("Simulado","Simulado completo 100 q. (07h–12h)"), ("Simulado","Simulado (continuação)")) if sim_sun
                    else (("Português","Interpretação de textos FGV"), ("Revisão","Revisão da semana + painel"))}
    return table[wd]

days = []
d = plano.START
while d < plano.START + dt.timedelta(weeks=26):
    if d.weekday() in WEEKDAYS:
        n = (d - plano.START).days // 7 + 1
        days.append((d, WEEKDAYS[d.weekday()], n) + blocks(n, d.weekday()))
    d += dt.timedelta(days=1)
FIRST = 6; LAST = FIRST + len(days) - 1

# ================= 1. COMECE AQUI =================
ws = sheet("Comece aqui", True, "1D6B62")
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 26; ws.column_dimensions["C"].width = 92
ws["B1"] = "Planilha de estudos da Patrícia"; ws["B1"].font = title
ws["B2"] = "Perito Legista PCERJ · 2h por dia · 02/11/2026 a 02/05/2027"; ws["B2"].font = sub
ws["B4"] = "Regra única: você só preenche as células AMARELAS. O resto a planilha faz sozinha."; ws["B4"].font = Font(name=F, size=12, bold=True)
ws["B4"].fill = INP; ws.merge_cells("B4:C4"); ws.row_dimensions[4].height = 24
steps = [
 ("QUANDO", "O QUE FAZER", True),
 ("Todo dia (2 min)", "Aba DIÁRIO → vá até a linha de hoje (fica azul). Ela já diz o que estudar em cada bloco. Ao terminar, marque ✓ no bloco e escreva quantas questões fez e quantas acertou.", False),
 ("Terminou a teoria de um tópico", "Aba EDITAL → escreva a data na coluna 'Estudei em'. A planilha calcula sozinha quando revisar (7 e 30 dias).", False),
 ("Errou uma questão importante", "Aba ERROS → uma linha: o que errou e a regra certa em uma frase.", False),
 ("Domingo (10 min)", "Aba PAINEL → veja a cor da semana. Depois, nas abas EDITAL e ERROS, faça tudo que estiver 🟡. Mande um print do painel para a mentora.", False),
 ("Dia de simulado", "Aba SIMULADOS → escreva quantas acertou em cada matéria.", False),
 ("Quer ver a semana toda", "Aba O QUE ESTUDAR → o conteúdo das 26 semanas, de uma vez.", False),
]
r = 6
for a, b, h in steps:
    for col, v in ((2, a), (3, b)):
        x = ws.cell(row=r, column=col, value=v); x.border = box; x.alignment = W
        x.font = H if h else (bold if col == 2 else base)
        if h: x.fill = HF
    ws.row_dimensions[r].height = 20 if h else 34; r += 1
r += 1
ws.cell(row=r, column=2, value="O QUE AS CORES QUEREM DIZER").font = bold; r += 1
for sym, txt, f in (("🟢 No ritmo", "Está tudo certo. Continue.", GREEN), ("🟡 Atenção", "Tem revisão para fazer ou alguma matéria abaixo de 60%. Resolva no domingo.", YEL),
                    ("🔴 Alerta", "Semana muito abaixo do plano. Fale com a mentora.", RED), ("Linha azul", "É o dia de hoje no Diário.", TODAYF)):
    a = ws.cell(row=r, column=2, value=sym); a.fill = f; a.font = bold; a.border = box
    b = ws.cell(row=r, column=3, value=txt); b.font = base; b.border = box; r += 1
r += 1
ws.cell(row=r, column=2, value="Dúvidas sobre COMO estudar: consulte o Guia da Aluna (PDF).").font = sub

# ================= 2. DIÁRIO =================
ws = sheet("Diário", tab="F2C94C")
top(ws, "Diário · marque ✓ em cada bloco feito", "Os dias e o conteúdo de cada bloco já estão preenchidos. Você só marca ✓ e anota questões e acertos (células amarelas).")
ws["A3"] = "Hoje:"; ws["A3"].font = bold; ws["B3"] = "=TODAY()"; ws["B3"].number_format = "dd/mm/yyyy"; ws["B3"].font = bold; ws["B3"].fill = TODAYF
cols = ["Data","Dia","Semana","Bloco 1 · 07h–08h","Feito?","Questões","Acertos","Bloco 2 · 08h–09h","Feito?","Questões","Acertos","Anotações (opcional)","m1","m2"]
head(ws, 5, cols, [11,10,8,44,7,9,9,44,7,9,9,30,4,4])
for i, (day, wdname, n, (d1, t1), (d2, t2)) in enumerate(days):
    r = FIRST + i; zebra = WEEKA if n % 2 == 0 else None
    put(ws, r, 1, day, fmt="dd/mm/yy", f=zebra); put(ws, r, 2, wdname, f=zebra); put(ws, r, 3, f"S{n}", f=zebra)
    put(ws, r, 4, f"{d1} · {t1}" if d1 not in ("Revisão","Simulado","Livre") else t1, al=W, f=GREY if d1 == "Livre" else zebra)
    put(ws, r, 8, f"{d2} · {t2}" if d2 not in ("Revisão","Simulado","Livre") else t2, al=W, f=GREY if d2 == "Livre" else zebra)
    for c in (5,6,7,9,10,11): put(ws, r, c, f=INP)
    put(ws, r, 12, f=INP, al=W)
    put(ws, r, 13, d1, font=small); put(ws, r, 14, d2, font=small)
    ws.row_dimensions[r].height = 30
ws.column_dimensions["M"].hidden = True; ws.column_dimensions["N"].hidden = True
dv(ws, f"E{FIRST}:E{LAST} I{FIRST}:I{LAST}", '"✓"')
dv(ws, f"F{FIRST}:F{LAST} J{FIRST}:J{LAST}", "0", kind="whole", operator="greaterThanOrEqual", error="Digite um número inteiro.", showErrorMessage=True)
dv(ws, f"G{FIRST}:G{LAST}", f"=G{FIRST}<=F{FIRST}", kind="custom", error="Acertos não pode ser maior que o número de questões.", showErrorMessage=True)
dv(ws, f"K{FIRST}:K{LAST}", f"=K{FIRST}<=J{FIRST}", kind="custom", error="Acertos não pode ser maior que o número de questões.", showErrorMessage=True)
ws.conditional_formatting.add(f"A{FIRST}:D{LAST}", FormulaRule(formula=[f"$A{FIRST}=TODAY()"], fill=TODAYF, font=Font(name=F, bold=True)))
ws.conditional_formatting.add(f"H{FIRST}:H{LAST}", FormulaRule(formula=[f"$A{FIRST}=TODAY()"], fill=TODAYF, font=Font(name=F, bold=True)))
ws.conditional_formatting.add(f"E{FIRST}:E{LAST}", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
ws.conditional_formatting.add(f"I{FIRST}:I{LAST}", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
ws.freeze_panes = f"D{FIRST}"

# ================= 3. PAINEL =================
ws = sheet("Painel", tab="27AE60")
top(ws, "Painel · como foi a sua semana", "Tudo aqui é automático. Olhe a coluna SITUAÇÃO todo domingo e mande um print para a mentora.")
D = "Diário!"
rngC, rngE, rngI, rngF, rngG, rngJ, rngK, rngM, rngN = (f"{D}${c}${FIRST}:${c}${LAST}" for c in "CEIFGJKMN")
kpis = [("Blocos feitos", f"=COUNTIF({rngE},\"✓\")+COUNTIF({rngI},\"✓\")", "0"),
        ("Questões feitas", f"=SUM({rngF})+SUM({rngJ})", "0"),
        ("Acerto geral", f"=IF((SUM({rngF})+SUM({rngJ}))=0,\"–\",(SUM({rngG})+SUM({rngK}))/(SUM({rngF})+SUM({rngJ})))", "0%"),
        ("Edital estudado", f"=COUNTA(Edital!D6:D{5+len(plano.TOPICS)})/{len(plano.TOPICS)}", "0%"),
        ("Revisões para fazer", f"=COUNTIF(Edital!I6:I{5+len(plano.TOPICS)},\"🟡*\")+COUNTIF(Erros!H6:H405,\"🟡*\")", "0")]
for j, (lab, fml, fmt) in enumerate(kpis):
    col = 1 + j*2
    a = ws.cell(row=3, column=col, value=lab); a.font = small; a.alignment = C
    b = ws.cell(row=4, column=col, value=fml); b.font = Font(name=F, size=16, bold=True, color=TEAL); b.alignment = C; b.number_format = fmt
    ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+1); ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col+1)
ws.row_dimensions[4].height = 30
cols = ["Semana","Datas","Blocos planejados","Blocos feitos","% feito","Questões","Meta de questões","Acerto Medicina Legal","Acerto Português","Acerto Direito","SITUAÇÃO"]
head(ws, 6, cols, [9,15,11,10,9,10,10,12,12,12,52])
for i, w in enumerate(plano.WEEKS):
    r = 7+i; n = w["n"]; S = f'"S{n}"'
    put(ws, r, 1, f"S{n}"); put(ws, r, 2, w["datas"])
    put(ws, r, 3, f'=COUNTIFS({rngC},{S},{rngM},"<>Livre")+COUNTIFS({rngC},{S},{rngN},"<>Livre")')
    put(ws, r, 4, f'=COUNTIFS({rngC},{S},{rngE},"✓")+COUNTIFS({rngC},{S},{rngI},"✓")')
    put(ws, r, 5, f'=IF(D{r}=0,"",D{r}/C{r})', fmt="0%")
    put(ws, r, 6, f'=SUMIFS({rngF},{rngC},{S})+SUMIFS({rngJ},{rngC},{S})')
    put(ws, r, 7, w["questoes"])
    for c, disc in zip((8,9,10), DISC):
        q = f'(SUMIFS({rngF},{rngC},{S},{rngM},"{disc}")+SUMIFS({rngJ},{rngC},{S},{rngN},"{disc}"))'
        a = f'(SUMIFS({rngG},{rngC},{S},{rngM},"{disc}")+SUMIFS({rngK},{rngC},{S},{rngN},"{disc}"))'
        put(ws, r, c, f'=IF({q}=0,"",{a}/{q})', fmt="0%")
    end = f"DATE({w['fim'].year},{w['fim'].month},{w['fim'].day})"
    det = (f'IF(E{r}<0.8,"fez poucos blocos; ","")&IF(AND(ISNUMBER(H{r}),H{r}<0.6),"reforçar Medicina Legal; ","")'
           f'&IF(AND(ISNUMBER(I{r}),I{r}<0.6),"reforçar Português; ","")&IF(AND(ISNUMBER(J{r}),J{r}<0.6),"reforçar Direito; ","")')
    put(ws, r, 11, f'=IF(D{r}=0,IF(TODAY()>{end},"🔴 Semana sem registro",""),IF(E{r}<0.5,"🔴 Fale com a mentora",IF({det}="","🟢 No ritmo","🟡 Atenção: "&{det})))', al=W)
ws.conditional_formatting.add("K7:K32", FormulaRule(formula=['LEFT(K7,1)="🟢"'], fill=GREEN))
ws.conditional_formatting.add("K7:K32", FormulaRule(formula=['LEFT(K7,1)="🟡"'], fill=YEL))
ws.conditional_formatting.add("K7:K32", FormulaRule(formula=['LEFT(K7,1)="🔴"'], fill=RED))
ws.conditional_formatting.add("H7:J32", FormulaRule(formula=['AND(ISNUMBER(H7),H7<0.6)'], font=Font(name=F, bold=True, color="9C1C1C")))
ch = BarChart(); ch.title = "Blocos planejados x feitos"; ch.height = 7; ch.width = 20
ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=6, max_row=32), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=7, max_row=32)); ws.add_chart(ch, "M6")
ws.freeze_panes = "C7"

# ================= 4. O QUE ESTUDAR =================
ws = sheet("O que estudar", tab="2F80ED")
top(ws, "O que estudar · as 26 semanas", "Visão geral. O Diário já mostra o conteúdo de cada dia. Marque ✓ quando fechar a semana.")
cols = ["Semana","Datas","Medicina Legal (4h)","Português (3h)","Direito (2h)","Meta de questões","Revisão de 30 dias","Destaque","Semana concluída?"]
head(ws, 5, cols, [8,15,38,32,38,12,12,30,11])
for i, w in enumerate(plano.WEEKS):
    r = 6+i; ml, pt, di = plano.week_short(w["n"])
    dest = w["simulado"] or ("Semana leve (" + plano.RECOVERY[w["n"]] + ")" if w["rec"] else ("Semana da prova: descansar" if w["n"] == 26 else ""))
    f_ = YEL if w["rec"] else None
    for c, v in enumerate([f"S{w['n']}", w["datas"], ml, pt, di, w["questoes"], w["rev30"], dest], 1):
        put(ws, r, c, v, al=W if c in (3,4,5,8) else C, f=f_)
    put(ws, r, 9, f=INP); ws.row_dimensions[r].height = 30
    if w["rec"] or w["n"] == 26: ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
dv(ws, "I6:I31", '"✓"'); ws.conditional_formatting.add("I6:I31", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
ws.freeze_panes = "C6"

# ================= 5. EDITAL =================
ws = sheet("Edital", tab="9B51E0")
top(ws, "Edital · os 51 tópicos e suas revisões", "Quando terminar a teoria de um tópico, escreva a data em 'Estudei em'. Quando fizer a revisão, marque ✓.")
cols = ["Semana","Matéria","Tópico","Estudei em","Revisar 7 dias em","Feita?","Revisar 30 dias em","Feita?","Situação","O que o edital pede (detalhe)"]
head(ws, 5, cols, [8,15,40,12,13,8,13,8,24,70])
topics = sorted(plano.TOPICS, key=lambda t: (t[2], DISC.index(t[0])))
for i, (disc, code, wk, full) in enumerate(topics):
    r = 6+i; fim = plano.wk_dates(wk)[1]; end = f"DATE({fim.year},{fim.month},{fim.day})"
    put(ws, r, 1, f"S{wk}"); put(ws, r, 2, disc); put(ws, r, 3, plano.SHORT[code], al=W)
    put(ws, r, 4, f=INP, fmt="dd/mm/yy")
    put(ws, r, 5, f'=IF(D{r}="","",D{r}+7)', fmt="dd/mm/yy"); put(ws, r, 6, f=INP)
    put(ws, r, 7, f'=IF(D{r}="","",D{r}+30)', fmt="dd/mm/yy"); put(ws, r, 8, f=INP)
    put(ws, r, 9, f'=IF(D{r}="",IF(TODAY()>{end},"🔴 Atrasado","⚪ A estudar"),IF(AND(F{r}<>"✓",TODAY()>=E{r}),"🟡 Fazer revisão de 7 dias",IF(AND(H{r}<>"✓",TODAY()>=G{r}),"🟡 Fazer revisão de 30 dias","🟢 Em dia")))')
    put(ws, r, 10, full, font=small, al=W); ws.row_dimensions[r].height = 30
L = 5+len(topics)
dv(ws, f"F6:F{L} H6:H{L}", '"✓"')
for c, f_ in (("🟢", GREEN), ("🟡", YEL), ("🔴", RED)):
    ws.conditional_formatting.add(f"I6:I{L}", FormulaRule(formula=[f'LEFT(I6,1)="{c}"'], fill=f_))
ws.conditional_formatting.add(f"F6:F{L}", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
ws.conditional_formatting.add(f"H6:H{L}", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
ws.freeze_panes = "D6"

# ================= 6. ERROS =================
ws = sheet("Erros", tab="EB5757")
top(ws, "Caderno de erros", "Uma linha por erro importante (ou acerto no chute). Revise na data indicada e marque ✓.")
cols = ["Data","Matéria","Assunto","O que eu errei","Regra certa (uma frase)","Revisar em","Revisei?","Situação"]
head(ws, 5, cols, [11,15,24,40,46,12,9,18])
ex = ["03/11/26","Medicina Legal","Documentos médico-legais","Achei que parecer e laudo eram a mesma coisa","Parecer é a opinião sobre uma perícia já feita, sem novo exame","10/11/26","✓","exemplo"]
for c, v in enumerate(ex, 1): put(ws, 4, c, v, font=Font(name=F, size=9, italic=True, color="888888"), f=GREY, al=W)
for r in range(6, 406):
    for c in range(1, 9):
        if c == 6: put(ws, r, c, f'=IF(A{r}="","",A{r}+7)', fmt="dd/mm/yy")
        elif c == 8: put(ws, r, c, f'=IF(A{r}="","",IF(G{r}="✓","🟢 Revisado",IF(TODAY()>=F{r},"🟡 Revisar","Aguardando")))')
        else: put(ws, r, c, f=INP, fmt="dd/mm/yy" if c == 1 else None, al=W if c in (3,4,5) else C)
dv(ws, "B6:B405", '"Medicina Legal,Português,Direito"'); dv(ws, "G6:G405", '"✓"')
ws.conditional_formatting.add("H6:H405", FormulaRule(formula=['LEFT(H6,1)="🟡"'], fill=YEL))
ws.conditional_formatting.add("H6:H405", FormulaRule(formula=['LEFT(H6,1)="🟢"'], fill=GREEN))
ws.freeze_panes = "A6"

# ================= 7. SIMULADOS =================
ws = sheet("Simulados", tab="F2994A")
top(ws, "Simulados", "Escreva só quantas questões acertou em cada matéria. Na prova é preciso acertar pelo menos 50% em CADA matéria.")
cols = ["Simulado","Semana","Data","Acertos ML","de","Acertos Port.","de","Acertos Dir.","de","% ML","% Port.","% Dir.","% Geral","Resultado"]
head(ws, 5, cols, [24,8,11,10,5,11,5,10,5,8,8,8,9,44])
for i, (nm, wk, a, b, c_) in enumerate([("Mini-simulado diagnóstico","S13",20,12,8),("Simulado 1","S21",50,30,20),("Simulado 2","S23",50,30,20),("Simulado 3","S25",50,30,20),("Extra","",50,30,20),("Extra","",50,30,20)]):
    r = 6+i
    put(ws, r, 1, nm, al=W); put(ws, r, 2, wk); put(ws, r, 3, f=INP, fmt="dd/mm/yy")
    put(ws, r, 4, f=INP); put(ws, r, 5, a, f=GREY); put(ws, r, 6, f=INP); put(ws, r, 7, b, f=GREY); put(ws, r, 8, f=INP); put(ws, r, 9, c_, f=GREY)
    put(ws, r, 10, f'=IF(D{r}="","",D{r}/E{r})', fmt="0%"); put(ws, r, 11, f'=IF(F{r}="","",F{r}/G{r})', fmt="0%"); put(ws, r, 12, f'=IF(H{r}="","",H{r}/I{r})', fmt="0%")
    put(ws, r, 13, f'=IF(OR(D{r}="",F{r}="",H{r}=""),"",(D{r}+F{r}+H{r})/(E{r}+G{r}+I{r}))', fmt="0%")
    low = f'IF(J{r}<0.5,"ML ","")&IF(K{r}<0.5,"Português ","")&IF(L{r}<0.5,"Direito ","")'
    put(ws, r, 14, f'=IF(M{r}="","",IF({low}<>"","🔴 Abaixo de 50% em: "&{low},IF(M{r}>=0.7,"🟢 Meta de 70% atingida","🟡 Passa no mínimo; a meta é 70%")))', al=W)
    for col, mx in ((4,"E"),(6,"G"),(8,"I")):
        L_ = ws.cell(row=r, column=col).column_letter
        dv(ws, f"{L_}{r}", f"={L_}{r}<={mx}{r}", kind="custom", error="Acertos maior que o total de questões.", showErrorMessage=True)
for c, f_ in (("🟢", GREEN), ("🟡", YEL), ("🔴", RED)):
    ws.conditional_formatting.add("N6:N11", FormulaRule(formula=[f'LEFT(N6,1)="{c}"'], fill=f_))

for s in wb.worksheets:
    s.sheet_properties.pageSetUpPr.fitToPage = True; s.page_setup.orientation = "landscape"; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
wb.save(OUT); print("saved", len(days), "dias")
