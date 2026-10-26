import sys, plano
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = "Arial"
TEAL = "1D6B62"; INPUT = "FFF4C2"; GREY = "EEF2F1"; EX = "F3F3F3"
hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", fgColor=TEAL)
inp = PatternFill("solid", fgColor=INPUT)
grey = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="C9D3D0"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=16, bold=True, color=TEAL); sub = Font(name=F, size=10, italic=True, color="5A6863")
N = len(plano.TOPICS); REG_LAST = 504; ERR_LAST = 404

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name; ws.sheet_view.showGridLines = False
    return ws
def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c); cell.font = hfont; cell.fill = hfill
        cell.alignment = center; cell.border = box
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
def put(ws, r, c, v, font=base, fill=None, align=wrap, fmt=None):
    cell = ws.cell(row=r, column=c, value=v); cell.font = font; cell.alignment = align; cell.border = box
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    return cell
def dv_list(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(dv); dv.add(rng)

# ---------- 0. Como usar ----------
ws = sheet("Como usar", first=True)
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 110
ws["B1"] = f"Planilha de estudos · {plano.ALUNA}"; ws["B1"].font = title
ws["B2"] = "Perito Legista PCERJ · 2h por dia · 26 semanas (02/11/2026 a 02/05/2027) · base: Edital nº 05/2021"; ws["B2"].font = sub
txt = [
 ("LEGENDA", None),
 ("Células AMARELAS = você preenche. Todo o resto é calculado sozinho (não apague fórmulas).", inp),
 ("", None),
 ("ROTINA DE USO", None),
 ("1. Todo dia, ao terminar o estudo (2 min): aba 'Registro diário' → uma linha por bloco (data, disciplina, código do tópico, atividade, minutos, questões, acertos).", None),
 ("2. Ao terminar a teoria de um tópico: aba 'Edital' → marque 'Teoria concluída = Sim' e a data do 1º estudo. As datas das revisões (24h, 7 e 30 dias) aparecem sozinhas.", None),
 ("3. Errou uma questão importante ou chutou e acertou: aba 'Caderno de erros' → uma linha. A data de revisão aparece sozinha.", None),
 ("4. Domingo, no bloco de revisão: abra 'Painel semanal' e leia a coluna ALERTA da semana. Faça as revisões marcadas como PENDENTE na aba 'Edital'. Marque o status da semana na aba 'Cronograma'.", None),
 ("5. Simulados: aba 'Simulados'. Teste físico: aba 'TAF'. Ao fim de cada dia: marque os blocos feitos na aba 'Checklist semanal'.", None),
 ("", None),
 ("ABAS", None),
 ("Grade semanal — horários fixos de cada dia (07h–09h).", None),
 ("Cronograma — o que estudar em cada uma das 26 semanas, metas de questões, revisões e simulados.", None),
 ("Checklist semanal — marque cada bloco feito, os flashcards e o ritual de domingo. Substitui a folha impressa.", None),
 ("Edital — os 51 tópicos do edital com checklist, datas de revisão, questões e % de acerto por tópico.", None),
 ("Registro diário — seu diário de estudo. É daqui que saem todos os números do painel.", None),
 ("Painel semanal — horas planejadas x executadas, questões, acerto por disciplina, % do edital e alertas.", None),
 ("Caderno de erros — registro e revisão dos erros.", None),
 ("Simulados — notas por módulo e checagem do mínimo de 50% em cada módulo.", None),
 ("TAF — evolução nos 4 testes físicos (índices femininos do edital de 2021).", None),
 ("", None),
 ("REGRAS DE AJUSTE (o painel avisa)", None),
 ("• Executou menos de 80% das horas planejadas por 2 semanas seguidas → o plano está acima da rotina real: converse com a mentora para reduzir, em vez de acumular.", None),
 ("• Uma disciplina com acerto abaixo de 60% → ela ganha 1 bloco por 2 semanas, tirado da disciplina com melhor acerto (Medicina Legal nunca perde bloco).", None),
 ("• Revisões pendentes há mais de 1 semana → o próximo bloco de domingo é só revisão.", None),
 ("", None),
 ("Os tópicos seguem o Anexo I do edital de 2021. Quando sair o novo edital, a mentora atualiza a aba 'Edital'.", None),
 ("Dúvidas sobre COMO estudar (blocos, revisões, simulados, atrasos, checklists): consulte o GUIA DA ALUNA (PDF). Esta planilha e o guia são os únicos materiais da mentoria.", None),
]
r = 4
for t, fill in txt:
    c = ws.cell(row=r, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = bold if t.isupper() and t else base
    if fill: c.fill = fill
    r += 1

# ---------- 1. Grade semanal ----------
ws = sheet("Grade semanal")
ws["A1"] = "Grade semanal · 2h por dia (2 blocos de 1h)"; ws["A1"].font = title
ws["A2"] = "Cada bloco começa com 10 min de revisão de 24h da mesma matéria (escrever de memória 5 pontos do conteúdo anterior)."; ws["A2"].font = sub
header(ws, 4, ["Dia","Bloco","Horário","Disciplina","Atividade (Fase 1 · S1–S20)","Atividade (Fase 2 · S21–S25)"], [12,7,14,16,42,48])
fase2 = {("Segunda","1"):"Correção do simulado / questões ML",("Segunda","2"):"Questões ML",("Quarta","1"):"Questões Português",
         ("Quarta","2"):"Lei seca + questões Direito",("Quinta","1"):"Teoria dirigida aos erros de ML",("Quinta","2"):"Questões Português",
         ("Sexta","1"):"Questões ML",("Sexta","2"):"Questões Direito",("Domingo","1"):"Interpretação FGV (ou simulado 07h–12h)",
         ("Domingo","2"):"Revisão + painel (ou simulado)"}
for i, (d,b,h,disc,at) in enumerate(plano.GRADE):
    r = 5+i; off = disc == "Plantão"
    for c, v in enumerate([d,b,h,disc,at,fase2.get((d,b),"Sem estudo" if off else "")],1):
        put(ws, r, c, v, fill=grey if off else None, align=center if c<=3 else wrap)
r = 5+len(plano.GRADE)+1
for line in ["Total: Medicina Legal 4h · Português 3h · Direito 2h · Revisão 1h = 10h por semana.",
             "Semanas de simulado (S21, S23, S25): o domingo vira simulado de 100 questões das 07h às 12h (duração real da prova).",
             "Semanas de recuperação (S8 Natal, S15 Carnaval): carga reduzida (≈5h), só revisão e atrasos."]:
    ws.cell(row=r, column=1, value=line).font = base; r += 1
ws.freeze_panes = "A5"

# ---------- 2. Cronograma ----------
ws = sheet("Cronograma")
ws["A1"] = "Cronograma · 26 semanas"; ws["A1"].font = title
ws["A2"] = "Status e Observações são seus (amarelo). Os códigos (ML01, PT01, DI01…) são os mesmos da aba Edital."; ws["A2"].font = sub
cols = ["Semana","Datas","Fase","Medicina Legal","Português","Direito","Meta de questões","Horas plan.","Revisão 30 dias","Simulado / TAF","Resultado esperado","Status","Observações"]
header(ws, 4, cols, [8,15,14,44,34,40,22,9,10,26,30,15,28])
for i, w in enumerate(plano.WEEKS):
    r = 5+i
    extra = "; ".join(x for x in [w["simulado"], "Teste físico (TAF)" if w["taf"] else ""] if x)
    vals = [f"S{w['n']}", w["datas"], w["fase"],
            w["resumo"] if w["rec"] or w["n"]==26 else w["ml"], "" if w["rec"] or w["n"]==26 else w["pt"], "" if w["rec"] or w["n"]==26 else w["di"],
            f"{w['questoes']} ({w['q_split']})", w["horas"], w["rev30"], extra, w["esperado"], "", ""]
    for c, v in enumerate(vals, 1):
        put(ws, r, c, v, fill=inp if c>=12 else (PatternFill("solid", fgColor="FBF0E3") if w["rec"] else None),
            align=center if c in (1,3,8,9) else wrap)
    if w["rec"] or w["n"]==26: ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
dv_list(ws, '"Não iniciada,Em andamento,Concluída,Atrasada"', "L5:L30")
put(ws, 32, 1, "Total de horas planejadas:", font=bold, align=Alignment(horizontal="left")); ws.merge_cells("A32:G32")
put(ws, 32, 8, "=SUM(H5:H30)", font=bold, align=center)
ws.conditional_formatting.add("L5:L30", CellIsRule(operator="equal", formula=['"Concluída"'], fill=PatternFill("solid", fgColor="CDEBD9")))
ws.conditional_formatting.add("L5:L30", CellIsRule(operator="equal", formula=['"Atrasada"'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.freeze_panes = "D5"


# ---------- 2b. Checklist semanal ----------
ws = sheet("Checklist semanal")
ws["A1"] = "Checklist semanal · marque ✓ em cada bloco concluído"; ws["A1"].font = title
ws["A2"] = "Bloco 1 = 07h–08h · Bloco 2 = 08h–09h. Terça e sábado são plantão. Na semana de simulado, os 2 blocos de domingo = simulado."; ws["A2"].font = sub
cols = ["Semana","Datas","Seg B1","Seg B2","Qua B1","Qua B2","Qui B1","Qui B2","Sex B1","Sex B2","Dom B1","Dom B2","Blocos plan.","Blocos feitos","% feito","Dias com flashcards (0–5)","Ritual de domingo feito?","Print enviado à mentora?"]
header(ws, 4, cols, [8,15,7,7,7,7,7,7,7,7,7,7,9,9,8,12,12,12])
for i, w in enumerate(plano.WEEKS):
    r = 5+i
    put(ws, r, 1, f"S{w['n']}", align=center); put(ws, r, 2, w["datas"], align=center)
    for c in range(3, 13): put(ws, r, c, None, fill=inp, align=center)
    put(ws, r, 13, 5 if w["rec"] else (6 if w["n"]==26 else 10), align=center)
    put(ws, r, 14, f'=COUNTIF(C{r}:L{r},"✓")', align=center)
    put(ws, r, 15, f'=IF(N{r}=0,"",N{r}/M{r})', fmt="0%", align=center)
    for c in (16,17,18): put(ws, r, c, None, fill=inp, align=center)
dv_list(ws, '"✓"', "C5:L30"); dv_list(ws, '"0,1,2,3,4,5"', "P5:P30"); dv_list(ws, '"Sim,Não"', "Q5:R30")
ws.conditional_formatting.add("C5:L30", CellIsRule(operator="equal", formula=['"✓"'], fill=PatternFill("solid", fgColor="CDEBD9")))
ws.conditional_formatting.add("O5:O30", FormulaRule(formula=['AND(ISNUMBER(O5),O5<0.8)'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.freeze_panes = "C5"

# ---------- 3. Edital ----------
ws = sheet("Edital")
ws["A1"] = "Checklist do edital · 51 tópicos"; ws["A1"].font = title
ws["A2"] = "Preencha só as colunas amarelas. Revisões vencidas aparecem como PENDENTE."; ws["A2"].font = sub
# resumo
put(ws, 3, 1, "Edital concluído", font=bold); put(ws, 3, 2, f'=COUNTIF(O7:O{6+N},"Concluído")/{N}', font=bold, fmt="0%", align=center)
put(ws, 3, 3, "ML", font=bold, align=center); put(ws, 3, 4, f'=COUNTIFS(A7:A{6+N},"Medicina Legal",O7:O{6+N},"Concluído")/COUNTIF(A7:A{6+N},"Medicina Legal")', fmt="0%", align=center)
put(ws, 3, 5, "Port.", font=bold, align=center); put(ws, 3, 6, f'=COUNTIFS(A7:A{6+N},"Português",O7:O{6+N},"Concluído")/COUNTIF(A7:A{6+N},"Português")', fmt="0%", align=center)
put(ws, 3, 7, "Dir.", font=bold, align=center); put(ws, 3, 8, f'=COUNTIFS(A7:A{6+N},"Direito",O7:O{6+N},"Concluído")/COUNTIF(A7:A{6+N},"Direito")', fmt="0%", align=center)
put(ws, 4, 1, "Revisões pendentes", font=bold); put(ws, 4, 2, f'=COUNTIF(P7:P{6+N},"PENDENTE")', font=bold, align=center)
put(ws, 4, 3, "Tópicos atrasados", font=bold); ws.merge_cells("C4:D4"); put(ws, 4, 5, f'=COUNTIF(O7:O{6+N},"Atrasado")', font=bold, align=center)
cols = ["Disciplina","Código","Tópico do edital","Semana prevista","Teoria concluída?","Data do 1º estudo","Rev. 24h","Rev. 7 dias","Rev. 30 dias","Rev. 7d feita?","Rev. 30d feita?","Questões","Acertos","% acerto","Status","Revisão"]
header(ws, 6, cols, [15,8,62,9,11,12,11,11,11,10,10,9,9,9,13,11])
for i, (disc, code, wk, t) in enumerate(plano.TOPICS):
    r = 7+i
    put(ws, r, 1, disc); put(ws, r, 2, code, align=center); put(ws, r, 3, t); put(ws, r, 4, wk, align=center)
    put(ws, r, 5, None, fill=inp, align=center); put(ws, r, 6, None, fill=inp, fmt="dd/mm/yy", align=center)
    put(ws, r, 7, f'=IF(F{r}="","",F{r}+1)', fmt="dd/mm/yy", align=center)
    put(ws, r, 8, f'=IF(F{r}="","",F{r}+7)', fmt="dd/mm/yy", align=center)
    put(ws, r, 9, f'=IF(F{r}="","",F{r}+30)', fmt="dd/mm/yy", align=center)
    put(ws, r, 10, None, fill=inp, align=center); put(ws, r, 11, None, fill=inp, align=center)
    put(ws, r, 12, f"=SUMIFS('Registro diário'!$G$6:$G${REG_LAST},'Registro diário'!$D$6:$D${REG_LAST},B{r})", align=center)
    put(ws, r, 13, f"=SUMIFS('Registro diário'!$H$6:$H${REG_LAST},'Registro diário'!$D$6:$D${REG_LAST},B{r})", align=center)
    put(ws, r, 14, f'=IF(L{r}=0,"",M{r}/L{r})', fmt="0%", align=center)
    put(ws, r, 15, f'=IF(AND(E{r}="Sim",L{r}>0),"Concluído",IF(OR(E{r}="Sim",L{r}>0),"Em andamento",IF(TODAY()>DATE(2026,11,2)+7*D{r}-1,"Atrasado","A fazer")))', align=center)
    put(ws, r, 16, f'=IF(F{r}="","",IF(OR(AND(J{r}<>"Sim",TODAY()>=H{r}),AND(K{r}<>"Sim",TODAY()>=I{r})),"PENDENTE","ok"))', align=center)
last = 6+N
dv_list(ws, '"Sim,Não"', f"E7:E{last}"); dv_list(ws, '"Sim,Não"', f"J7:K{last}")
ws.conditional_formatting.add(f"O7:O{last}", CellIsRule(operator="equal", formula=['"Concluído"'], fill=PatternFill("solid", fgColor="CDEBD9")))
ws.conditional_formatting.add(f"O7:O{last}", CellIsRule(operator="equal", formula=['"Atrasado"'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.conditional_formatting.add(f"P7:P{last}", CellIsRule(operator="equal", formula=['"PENDENTE"'], fill=PatternFill("solid", fgColor="F6D3CF"), font=Font(name=F, bold=True, color="9C1C1C")))
ws.conditional_formatting.add(f"N7:N{last}", FormulaRule(formula=[f'AND(ISNUMBER(N7),N7<0.6)'], font=Font(name=F, bold=True, color="9C1C1C")))
ws.freeze_panes = "D7"

# ---------- 4. Registro diário ----------
ws = sheet("Registro diário")
ws["A1"] = "Registro diário · uma linha por bloco de estudo"; ws["A1"].font = title
ws["A2"] = "A linha cinza é só exemplo (não entra nas contas). Preencha a partir da linha 6. Semana e % são automáticos."; ws["A2"].font = sub
cols = ["Data","Semana","Disciplina","Código do tópico","Atividade","Minutos","Questões","Acertos","% acerto","Observação"]
header(ws, 5, cols, [12,9,16,14,16,10,10,10,10,40])
exrow = ["03/11/2026 (ex.)", "1", "Medicina Legal", "ML01", "Questões", 60, 15, 10, "67%", "Exemplo: errei 2 sobre tipos de documento médico-legal"]
for c, v in enumerate(exrow, 1): put(ws, 4, c, v, font=Font(name=F, size=10, italic=True, color="777777"), fill=PatternFill("solid", fgColor=EX), align=center if c<10 else wrap)
ws.cell(row=4, column=1).comment = Comment("Linha de exemplo: mostra o formato. Os seus registros começam na linha 6.", "Mentoria")
for r in range(6, REG_LAST+1):
    put(ws, r, 1, None, fill=inp, fmt="dd/mm/yy", align=center)
    put(ws, r, 2, f'=IF(A{r}="","",INT((A{r}-DATE(2026,11,2))/7)+1)', align=center)
    for c in (3,4,5,6,7,8): put(ws, r, c, None, fill=inp, align=center)
    put(ws, r, 9, f'=IF(OR(G{r}="",G{r}=0),"",H{r}/G{r})', fmt="0%", align=center)
    put(ws, r, 10, None, fill=inp)
codes = ",".join(c for _,c,_,_ in plano.TOPICS)
# lista de códigos numa coluna oculta (lista >255 chars não cabe na validação direta)
for i, (_,c,_,_) in enumerate(plano.TOPICS): ws.cell(row=6+i, column=26, value=c).font = base
ws.cell(row=6+N, column=26, value="GERAL").font = base
ws.column_dimensions["Z"].hidden = True
dv_list(ws, '"Medicina Legal,Português,Direito,Revisão"', f"C6:C{REG_LAST}")
dv_list(ws, f"=$Z$6:$Z${6+N}", f"D6:D{REG_LAST}")
dv_list(ws, '"Teoria,Questões,Revisão,Lei seca,Simulado,Flashcards"', f"E6:E{REG_LAST}")
ws.cell(row=3, column=1, value="Código GERAL = revisão ou simulado que mistura vários tópicos.").font = sub
ws.freeze_panes = "A6"

# ---------- 5. Painel semanal ----------
ws = sheet("Painel semanal")
ws["A1"] = "Painel semanal"; ws["A1"].font = title
ws["A2"] = "Tudo aqui é automático (vem do Registro diário e do Edital). Leia a coluna ALERTA todo domingo."; ws["A2"].font = sub
put(ws, 3, 1, "Edital concluído hoje", font=bold); ws.merge_cells("A3:B3"); put(ws, 3, 3, "=Edital!B3", font=bold, fmt="0%", align=center)
put(ws, 3, 4, "Revisões pendentes", font=bold); ws.merge_cells("D3:E3"); put(ws, 3, 6, "=Edital!B4", font=bold, align=center)
put(ws, 3, 7, "Horas totais", font=bold); put(ws, 3, 8, "=SUM(D7:D32)", font=bold, fmt="0.0", align=center)
put(ws, 3, 9, "Questões totais", font=bold); ws.merge_cells("I3:J3"); put(ws, 3, 11, "=SUM(G7:G32)", font=bold, align=center)
cols = ["Semana","Datas","Horas plan.","Horas exec.","% execução","Meta questões","Questões feitas","Acerto geral","Acerto ML","Acerto Port.","Acerto Dir.","Edital previsto (acum.)","ALERTA"]
header(ws, 6, cols, [8,15,9,9,10,10,10,10,10,10,10,12,46])
R = "'Registro diário'!"
for i, w in enumerate(plano.WEEKS):
    r = 7+i; n = w["n"]
    put(ws, r, 1, n, align=center); put(ws, r, 2, w["datas"], align=center); put(ws, r, 3, w["horas"], align=center)
    put(ws, r, 4, f"=SUMIFS({R}$F$6:$F${REG_LAST},{R}$B$6:$B${REG_LAST},A{r})/60", fmt="0.0", align=center)
    put(ws, r, 5, f'=IF(D{r}=0,"",D{r}/C{r})', fmt="0%", align=center)
    put(ws, r, 6, w["questoes"], align=center)
    put(ws, r, 7, f"=SUMIFS({R}$G$6:$G${REG_LAST},{R}$B$6:$B${REG_LAST},A{r})", align=center)
    put(ws, r, 8, f'=IF(G{r}=0,"",SUMIFS({R}$H$6:$H${REG_LAST},{R}$B$6:$B${REG_LAST},A{r})/G{r})', fmt="0%", align=center)
    for c, disc in ((9,"Medicina Legal"),(10,"Português"),(11,"Direito")):
        L = ws.cell(row=1, column=c).column_letter
        q = f'SUMIFS({R}$G$6:$G${REG_LAST},{R}$B$6:$B${REG_LAST},$A{r},{R}$C$6:$C${REG_LAST},"{disc}")'
        a = f'SUMIFS({R}$H$6:$H${REG_LAST},{R}$B$6:$B${REG_LAST},$A{r},{R}$C$6:$C${REG_LAST},"{disc}")'
        put(ws, r, c, f'=IF({q}=0,"",{a}/{q})', fmt="0%", align=center)
    put(ws, r, 12, f'=COUNTIF(Edital!$D$7:$D${6+N},"<="&A{r})/{N}', fmt="0%", align=center)
    alert = (f'IF(AND(ISNUMBER(E{r}),E{r}<0.8),"Horas abaixo de 80%. ","")'
             f'&IF(AND(ISNUMBER(I{r}),I{r}<0.6),"ML abaixo de 60%: +1 bloco. ","")'
             f'&IF(AND(ISNUMBER(J{r}),J{r}<0.6),"Português abaixo de 60%: +1 bloco. ","")'
             f'&IF(AND(ISNUMBER(K{r}),K{r}<0.6),"Direito abaixo de 60%: +1 bloco. ","")'
             f'&IF(AND(G{r}>0,G{r}<0.8*F{r}),"Questões abaixo da meta. ","")')
    put(ws, r, 13, f'=IF(D{r}=0,"",IF({alert}="","Dentro do plano",{alert}))')
ws.conditional_formatting.add("M7:M32", FormulaRule(formula=['AND(M7<>"",M7<>"Dentro do plano")'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.conditional_formatting.add("M7:M32", CellIsRule(operator="equal", formula=['"Dentro do plano"'], fill=PatternFill("solid", fgColor="CDEBD9")))
ch = LineChart(); ch.title = "Acerto por disciplina (%)"; ch.height = 7.5; ch.width = 18; ch.y_axis.number_format = "0%"
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.x_axis.title = "Semana"
ch.add_data(Reference(ws, min_col=9, max_col=11, min_row=6, max_row=32), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=7, max_row=32)); ws.add_chart(ch, "O6")
bc = BarChart(); bc.title = "Horas planejadas x executadas"; bc.height = 7.5; bc.width = 18
bc.add_data(Reference(ws, min_col=3, max_col=4, min_row=6, max_row=32), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=1, min_row=7, max_row=32)); ws.add_chart(bc, "O22")
ws.freeze_panes = "C7"

# ---------- 6. Caderno de erros ----------
ws = sheet("Caderno de erros")
ws["A1"] = "Caderno de erros"; ws["A1"].font = title
ws["A2"] = "Uma linha por erro (ou acerto no chute). Revise na data indicada e marque Revisado = Sim."; ws["A2"].font = sub
put(ws, 3, 1, "Para revisar hoje", font=bold); put(ws, 3, 2, f'=COUNTIFS(H7:H{ERR_LAST},"<="&TODAY(),I7:I{ERR_LAST},"<>Sim")', font=bold, align=center)
for j, m in enumerate(["Não sabia","Confundi conceitos","Interpretação","Falta de atenção","Tempo"]):
    put(ws, 3, 3+j*2, m, font=bold, align=center); put(ws, 3, 4+j*2, f'=COUNTIF(E7:E{ERR_LAST},"{m}")', align=center)
cols = ["Data","Disciplina","Código do tópico","Questão / fonte","Motivo do erro","O que eu errei","Regra correta (1 linha)","Revisar em","Revisado?"]
header(ws, 6, cols, [12,16,13,24,18,40,46,12,11])
ex = ["03/11/2026 (ex.)","Medicina Legal","ML01","QConcursos nº 123456","Confundi conceitos","Achei que parecer era o mesmo que relatório","Parecer: opinião sobre perícia já feita, sem exame; relatório (laudo) descreve o exame","10/11/2026","Sim"]
for c, v in enumerate(ex, 1): put(ws, 5, c, v, font=Font(name=F, size=10, italic=True, color="777777"), fill=PatternFill("solid", fgColor=EX))
for r in range(7, ERR_LAST+1):
    for c in range(1,10):
        if c == 8: put(ws, r, c, f'=IF(A{r}="","",A{r}+7)', fmt="dd/mm/yy", align=center)
        else: put(ws, r, c, None, fill=inp, fmt="dd/mm/yy" if c==1 else None, align=center if c in (1,2,3,5,9) else wrap)
dv_list(ws, '"Medicina Legal,Português,Direito"', f"B7:B{ERR_LAST}")
dv_list(ws, "='Registro diário'!$Z$6:$Z$" + str(6+N), f"C7:C{ERR_LAST}")
dv_list(ws, '"Não sabia,Confundi conceitos,Interpretação,Falta de atenção,Tempo"', f"E7:E{ERR_LAST}")
dv_list(ws, '"Sim,Não"', f"I7:I{ERR_LAST}")
ws.conditional_formatting.add(f"H7:H{ERR_LAST}", FormulaRule(formula=[f'AND(H7<>"",H7<=TODAY(),I7<>"Sim")'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.freeze_panes = "A7"

# ---------- 7. Simulados ----------
ws = sheet("Simulados")
ws["A1"] = "Simulados"; ws["A1"].font = title
ws["A2"] = "Mínimo do edital: 50% em CADA módulo (ML 25/50 · Port 15/30 · Dir 10/20). Meta de segurança: acima de 70% no geral."; ws["A2"].font = sub
cols = ["Simulado","Semana","Data","Q. ML","Acertos ML","Q. Port.","Acertos Port.","Q. Dir.","Acertos Dir.","% ML","% Port.","% Dir.","% Total","Passa no mínimo?","Meta 70%?"]
header(ws, 4, cols, [26,8,12,7,10,8,11,7,10,8,8,8,9,13,10])
sims = [("Mini-simulado diagnóstico",13,20,12,8),("Simulado 1",21,50,30,20),("Simulado 2",23,50,30,20),("Simulado 3",25,50,30,20),
        ("Extra",None,50,30,20),("Extra",None,50,30,20)]
for i, (nm, wk, qm, qp, qd) in enumerate(sims):
    r = 5+i
    put(ws, r, 1, nm); put(ws, r, 2, wk, align=center); put(ws, r, 3, None, fill=inp, fmt="dd/mm/yy", align=center)
    put(ws, r, 4, qm, align=center); put(ws, r, 5, None, fill=inp, align=center)
    put(ws, r, 6, qp, align=center); put(ws, r, 7, None, fill=inp, align=center)
    put(ws, r, 8, qd, align=center); put(ws, r, 9, None, fill=inp, align=center)
    put(ws, r, 10, f'=IF(E{r}="","",E{r}/D{r})', fmt="0%", align=center)
    put(ws, r, 11, f'=IF(G{r}="","",G{r}/F{r})', fmt="0%", align=center)
    put(ws, r, 12, f'=IF(I{r}="","",I{r}/H{r})', fmt="0%", align=center)
    put(ws, r, 13, f'=IF(OR(E{r}="",G{r}="",I{r}=""),"",(E{r}+G{r}+I{r})/(D{r}+F{r}+H{r}))', fmt="0%", align=center)
    put(ws, r, 14, f'=IF(M{r}="","",IF(AND(J{r}>=0.5,K{r}>=0.5,L{r}>=0.5),"Sim","NÃO"))', align=center)
    put(ws, r, 15, f'=IF(M{r}="","",IF(M{r}>=0.7,"Sim","Ainda não"))', align=center)
ws.conditional_formatting.add("N5:N10", CellIsRule(operator="equal", formula=['"NÃO"'], fill=PatternFill("solid", fgColor="F6D3CF")))
put(ws, 12, 1, "Depois de cada simulado: corrija na segunda (bloco 1), lance os erros no Caderno de erros e anote os 3 tópicos com mais erros na aba Cronograma (Observações).", align=Alignment(wrap_text=False))

# ---------- 8. TAF ----------
ws = sheet("TAF")
ws["A1"] = "Teste de Aptidão Física (índices femininos · edital 2021)"; ws["A1"].font = title
ws["A2"] = "Cada teste tem uma única tentativa na prova. Teste a cada 4 semanas e registre aqui."; ws["A2"].font = sub
cols = ["Semana","Data","Flexão de braço (rep. em 1 min)","Abdominal (rep. em 1 min)","100 m (segundos)","1.600 m (minutos, ex.: 11,5)","Flexão","Abdominal","100 m","1.600 m"]
header(ws, 4, cols, [9,12,16,16,13,17,11,11,11,11])
put(ws, 5, 1, "Índice", font=bold, align=center); put(ws, 5, 2, "")
for c, v in zip(range(3,7), ["≥ 15","≥ 20","≤ 22","≤ 12"]): put(ws, 5, c, v, font=bold, align=center, fill=grey)
for i, wk in enumerate([4,8,12,16,20,24,None,None]):
    r = 6+i
    put(ws, r, 1, f"S{wk}" if wk else "Extra", align=center); put(ws, r, 2, None, fill=inp, fmt="dd/mm/yy", align=center)
    for c in range(3,7): put(ws, r, c, None, fill=inp, align=center)
    put(ws, r, 7, f'=IF(C{r}="","",IF(C{r}>=15,"Apto","Abaixo"))', align=center)
    put(ws, r, 8, f'=IF(D{r}="","",IF(D{r}>=20,"Apto","Abaixo"))', align=center)
    put(ws, r, 9, f'=IF(E{r}="","",IF(E{r}<=22,"Apto","Abaixo"))', align=center)
    put(ws, r, 10, f'=IF(F{r}="","",IF(F{r}<=12,"Apto","Abaixo"))', align=center)
ws.conditional_formatting.add("G6:J13", CellIsRule(operator="equal", formula=['"Abaixo"'], fill=PatternFill("solid", fgColor="F6D3CF")))
ws.conditional_formatting.add("G6:J13", CellIsRule(operator="equal", formula=['"Apto"'], fill=PatternFill("solid", fgColor="CDEBD9")))

for s in wb.worksheets:
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.orientation = "landscape"; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
wb.save(OUT); print("saved", OUT)
